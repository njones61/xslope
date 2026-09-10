#!/usr/bin/env python3
"""Generate the shipped documentation index the Studio assistant searches.

A pip install carries no ``docs/`` tree, so the assistant that answers "does
xslope support line loads?" has never seen the page that says it does.  It
answered NO, and offered a workaround, on a capability the template, the loader
and three documentation pages all carry.  This tool folds the documentation into
one resource that ships in the wheel — ``xslope/resources/docs_index.json`` —
so the answer comes from the installed version's own pages rather than from
whatever the model remembers.

What the index holds:

  * every page in ``mkdocs.yml``'s nav — its site-relative URL in the same form
    the assistant brief's page table uses (``lem/overview/``), its title, and
    every heading with the anchor the built site serves and the section's body
    text;
  * the input template — every sheet, its header row, and the label / help-box
    strings on it (cell text, the literals inside a help formula, and the text
    of the drawing shapes on the reinforce and piles sheets);
  * a capabilities map: each template sheet against the documentation sections
    that name it, so "where is this input documented" is a lookup rather than a
    search.

Section bodies are stored as one lzma+base64 blob with a per-section (offset,
length) into the decompressed text: 2.4 MB of markdown does not fit a resource
budget uncompressed, and compressing the sections one at a time compresses
almost nothing.  lzma rather than zlib because the choice is worth 260 KB in the
wheel on identical stdlib-only code - 2.4 MB of markdown packs to 0.6 MB against
zlib's 0.8 MB, and base64 charges a third on top of whichever it is.

Both this file and :mod:`tools.make_corpus_index` slugify headings the way
python-markdown's toc extension does; the anchor helpers are imported from there
so the two indexes can never disagree about where a section lives.

Everything is sorted and nothing records a timestamp, so a rebuild is
byte-identical and ``run_tests.py`` can regenerate the index in memory and fail
when the committed copy is stale (the ``docs_index_sync`` row).

Usage::

    python tools/build_docs_index.py            # rewrite the resource
    python tools/build_docs_index.py --check    # report staleness, write nothing
"""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import sys
import lzma
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools.make_corpus_index import (  # noqa: E402
    slugify, strip_inline_markup, unique_slug)

DOCS = REPO_ROOT / 'docs'
MKDOCS = REPO_ROOT / 'mkdocs.yml'
TEMPLATE = REPO_ROOT / 'docs' / 'inputs' / 'input_template.xlsx'
INDEX_PATH = REPO_ROOT / 'xslope' / 'resources' / 'docs_index.json'

SCHEMA_VERSION = 1

HEADING_RE = re.compile(r'^(#{1,6})\s+(.*?)\s*$')
EXPLICIT_ANCHOR_RE = re.compile(r'\s*\{#([A-Za-z0-9_\-]+)\}\s*$')
COMMENT_RE = re.compile(r'<!--.*?-->', re.S)
IMAGE_RE = re.compile(r'!\[[^\]]*\]\([^)]*\)')
HTML_IMG_RE = re.compile(r'<img\b[^>]*>', re.I)
ADMONITION_RE = re.compile(r'^(\s*)(?:!!!|\?\?\?\+?)\s+[\w\-]+(?:\s+"([^"]*)")?\s*$')
FORMULA_STRING_RE = re.compile(r'"([^"]{4,})"')
BLANKS_RE = re.compile(r'\n{3,}')


# --------------------------------------------------------------------------- #
# mkdocs nav — the page list and the URL structure
# --------------------------------------------------------------------------- #

def read_nav():
    """``[(md_path, nav_title)]`` for every page in ``mkdocs.yml``'s nav, in order.

    mkdocs.yml carries python-object tags mkdocs itself resolves (``!!python/name``
    style hooks are configured elsewhere, but plugin blocks still use YAML that
    ``safe_load`` accepts), so the file is read with the safe loader and only the
    ``nav`` key is used.
    """
    import yaml

    with open(MKDOCS, encoding='utf-8') as fh:
        cfg = yaml.safe_load(fh)
    pages = []

    def walk(node, title=None):
        if isinstance(node, str):
            if node.endswith('.md'):
                pages.append((node, title))
        elif isinstance(node, list):
            for item in node:
                walk(item)
        elif isinstance(node, dict):
            for key, value in node.items():
                walk(value, str(key))

    walk(cfg.get('nav') or [])
    return pages


def page_url(rel_md: str) -> str:
    """``lem/overview.md`` -> ``lem/overview/`` (the mkdocs pretty URL)."""
    parts = list(Path(rel_md).with_suffix('').parts)
    if parts and parts[-1] == 'index':
        parts = parts[:-1]
    return '/'.join(parts) + ('/' if parts else '')


# --------------------------------------------------------------------------- #
# Markdown -> sections
# --------------------------------------------------------------------------- #

def clean_body(lines) -> str:
    """A section body as prose to search: markup that carries no words removed.

    Dropped: HTML comments (which is also where the test tags live), images in
    both markdown and HTML spelling, and the ``!!!``/``???`` admonition marker —
    whose title, where it has one, is kept as the line's text because it is
    usually the only statement of what the box is about.  Tables and fenced code
    survive intact: a table of columns is exactly what a "where is this input"
    question needs back.
    """
    text = '\n'.join(lines)
    text = COMMENT_RE.sub('', text)
    out = []
    in_fence = False
    for line in text.split('\n'):
        stripped = line.lstrip()
        if stripped.startswith('```') or stripped.startswith('~~~'):
            in_fence = not in_fence
            out.append(line)
            continue
        if in_fence:
            out.append(line)
            continue
        m = ADMONITION_RE.match(line)
        if m:
            out.append(m.group(1) + (m.group(2) or ''))
            continue
        line = IMAGE_RE.sub('', line)
        line = HTML_IMG_RE.sub('', line)
        out.append(line.rstrip())
    text = '\n'.join(out)
    text = BLANKS_RE.sub('\n\n', text)
    return text.strip()


def parse_page(md_path: Path, nav_title):
    """``{url, path, title, sections}`` for one page.

    A section runs from its heading to the next heading of any level; anything
    above the first heading is the page's lead and is recorded under the empty
    anchor, so a page whose opening paragraph carries the answer is still
    reachable.
    """
    raw = md_path.read_text(encoding='utf-8')
    lines = raw.splitlines()

    heads = []          # (line_index, level, title, anchor)
    used = set()
    in_fence = False
    for i, line in enumerate(lines):
        stripped = line.lstrip()
        if stripped.startswith('```') or stripped.startswith('~~~'):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        m = HEADING_RE.match(line)
        if not m:
            continue
        rawtitle = m.group(2)
        explicit = EXPLICIT_ANCHOR_RE.search(rawtitle)
        title = strip_inline_markup(rawtitle)
        if explicit:
            anchor = explicit.group(1)
            used.add(anchor)
        else:
            anchor = unique_slug(slugify(title), used)
        heads.append((i, len(m.group(1)), title, anchor))

    sections = []
    lead = clean_body(lines[:heads[0][0]] if heads else lines)
    if lead:
        sections.append({'anchor': '', 'heading': '', 'level': 0, 'text': lead})
    for k, (line_no, level, title, anchor) in enumerate(heads):
        end = heads[k + 1][0] if k + 1 < len(heads) else len(lines)
        body = clean_body(lines[line_no + 1:end])
        sections.append({'anchor': anchor, 'heading': title, 'level': level,
                         'text': body})

    title = nav_title
    h1 = next((h[2] for h in heads if h[1] == 1), None)
    if not title:
        title = h1 or md_path.stem
    return {'url': page_url(str(md_path.relative_to(DOCS))),
            'path': str(md_path.relative_to(REPO_ROOT)).replace(os.sep, '/'),
            'title': title, 'sections': sections}


# --------------------------------------------------------------------------- #
# The input template — sheets, headers, help-box strings
# --------------------------------------------------------------------------- #

def _drawing_texts(xlsx_path: Path):
    """``{sheet_name: [text, ...]}`` for the shapes drawn on a sheet.

    The reinforce and piles sheets state every column's meaning in a drawn text
    box rather than in cells, so the vocabulary a reader searches for ("pullout",
    "tieback", "moment capacity") lives only in the drawing XML.  Read straight
    out of the zip: openpyxl does not keep shape text, and the template is never
    written by this tool.
    """
    import zipfile
    from xml.etree import ElementTree as ET

    ns_r = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    ns_a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    out = {}
    with zipfile.ZipFile(xlsx_path) as zf:
        names = set(zf.namelist())
        # workbook sheet order -> sheetN.xml, via the workbook rels
        wb = ET.fromstring(zf.read('xl/workbook.xml'))
        rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
        target = {r.get('Id'): r.get('Target') for r in rels}
        for sheet in wb.iter():
            if not sheet.tag.endswith('}sheet'):
                continue
            name = sheet.get('name')
            rid = sheet.get(ns_r + 'id')
            part = target.get(rid)
            if not name or not part:
                continue
            part = 'xl/' + part.lstrip('/').replace('worksheets/../', '')
            relpath = (os.path.dirname(part) + '/_rels/'
                       + os.path.basename(part) + '.rels')
            if relpath not in names:
                continue
            sheet_rels = ET.fromstring(zf.read(relpath))
            texts = []
            for rel in sheet_rels:
                tgt = rel.get('Target') or ''
                if 'drawing' not in tgt:
                    continue
                dpath = os.path.normpath(os.path.join(os.path.dirname(part), tgt))
                dpath = dpath.replace(os.sep, '/')
                if dpath not in names:
                    continue
                drawing = ET.fromstring(zf.read(dpath))
                for para in drawing.iter(ns_a + 'p'):
                    runs = [t.text or '' for t in para.iter(ns_a + 't')]
                    line = ''.join(runs).strip()
                    if line:
                        texts.append(line)
            if texts:
                out[name] = texts
    return out


def read_template(xlsx_path: Path = TEMPLATE):
    """``[{name, header_row, headers, strings}]`` — one record per template sheet.

    The shipped template is BLANK, so every string on a sheet is a label, an
    option word or a help legend; that is the whole searchable vocabulary of the
    input format and it is taken verbatim.  A help note written as a formula
    (the dloads sheet states its water-load rule that way) contributes the
    literals inside the formula.

    ``header_row`` is the first row within the top eight whose cells read as a
    row of short column labels; ``headers`` is that row.  It is what a query for
    a column name ("Tmax", "Angle") matches, not a parsing contract — the loader
    owns that.
    """
    import openpyxl

    drawings = _drawing_texts(xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 60)):
            cells = []
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                if isinstance(value, str) and value.startswith('='):
                    cells += [s.strip() for s in FORMULA_STRING_RE.findall(value)]
                elif isinstance(value, str):
                    if value.strip():
                        cells.append(value.strip())
            rows.append(cells)

        header_row, headers = None, []
        for i, cells in enumerate(rows[:8], start=1):
            short = [c for c in cells if len(c) <= 14]
            if len(short) >= 3 and len(short) == len(cells):
                header_row, headers = i, short
                break

        strings = []
        seen = set()
        for i, cells in enumerate(rows, start=1):
            if i == header_row:
                continue
            for c in cells:
                if c not in seen:
                    seen.add(c)
                    strings.append(c)
        for text in drawings.get(name, []):
            if text not in seen:
                seen.add(text)
                strings.append(text)
        sheets.append({'name': name, 'header_row': header_row,
                       'headers': headers, 'strings': strings})
    return sheets


# --------------------------------------------------------------------------- #
# Capabilities — which sections document which sheet
# --------------------------------------------------------------------------- #

def sheet_mention_re(sheet: str):
    """A pattern that matches a sheet NAMED, not a word that happens to match.

    Half the sheet names are ordinary words (``main``, ``piles``, ``circles``,
    ``profile``), so a bare substring search would attribute most of the
    documentation to most of the sheets.  A mention counts when the name is
    quoted, in backticks, bold, or written beside the word
    "sheet"/"worksheet"/"tab" — the input-template page titles each of its
    sections ``Worksheet: lloads``, so the colon has to be allowed between them.
    """
    esc = re.escape(sheet)
    return re.compile(
        r'`%s`|\'%s\'|"%s"|\*\*%s\*\*'
        r'|\b%s\s+(?:sheet|worksheet|tab)\b'
        r'|\b(?:sheet|worksheet|tab)s?\s*:?\s+%s\b'
        % (esc, esc, esc, esc, esc, esc), re.I)


def build_capabilities(pages, sheets):
    """``[{sheet, sections: [{url, anchor, heading}]}]`` — sheet to documentation."""
    out = []
    for sheet in sheets:
        pattern = sheet_mention_re(sheet['name'])
        hits = []
        for page in pages:
            for sec in page['sections']:
                hay = sec['heading'] + '\n' + sec['text']
                if pattern.search(hay):
                    hits.append({'url': page['url'], 'anchor': sec['anchor'],
                                 'heading': sec['heading'] or page['title']})
        out.append({'sheet': sheet['name'], 'sections': hits})
    return out


# --------------------------------------------------------------------------- #
# Build / serialize
# --------------------------------------------------------------------------- #

def build_index(verbose=True):
    """The whole index, as the dict that is serialized to the resource."""
    pages = []
    for rel, nav_title in read_nav():
        md = DOCS / rel
        if not md.exists():
            if verbose:
                print('  missing page (in nav, not on disk): %s' % rel)
            continue
        pages.append(parse_page(md, nav_title))
    pages.sort(key=lambda p: p['url'])

    sheets = read_template()
    caps = build_capabilities(pages, sheets)

    # Section bodies out of the records and into one blob, so the resource
    # compresses as prose rather than as a few thousand short strings.
    blob_parts, cursor = [], 0
    for page in pages:
        for sec in page['sections']:
            text = sec.pop('text')
            sec['offset'] = cursor
            sec['length'] = len(text)
            blob_parts.append(text)
            cursor += len(text)
    blob = ''.join(blob_parts)
    packed = base64.b64encode(
        lzma.compress(blob.encode('utf-8'),
                      preset=9 | lzma.PRESET_EXTREME)).decode('ascii')

    if verbose:
        n_sec = sum(len(p['sections']) for p in pages)
        print('  %d pages, %d sections, %.1f MB of text -> %.0f KB packed'
              % (len(pages), n_sec, len(blob) / 1e6, len(packed) / 1024))

    return {'schema_version': SCHEMA_VERSION,
            'encoding': 'lzma+base64',
            'base_url': 'https://xslope.readthedocs.io/en/latest/',
            'bodies': packed,
            'pages': pages,
            'template': {'sheets': sheets},
            'capabilities': caps}


def serialize(index) -> str:
    return json.dumps(index, indent=1, sort_keys=True, ensure_ascii=False) + '\n'


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--check', action='store_true',
                    help='report staleness and write nothing')
    args = ap.parse_args(argv)

    print('building the documentation index...')
    text = serialize(build_index())
    if args.check:
        if not INDEX_PATH.exists():
            print('STALE: %s is missing' % INDEX_PATH)
            return 1
        if INDEX_PATH.read_text(encoding='utf-8') != text:
            print('STALE: %s differs from a rebuild' % INDEX_PATH)
            return 1
        print('current: %s' % INDEX_PATH)
        return 0
    INDEX_PATH.write_text(text, encoding='utf-8')
    print('wrote %s (%.0f KB)' % (INDEX_PATH, INDEX_PATH.stat().st_size / 1024))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
